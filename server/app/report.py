"""The research report as a file to share (contract §12, Export): PDF (reportlab) and PowerPoint
(python-pptx). Both carry the same content: the verdict, the design plan with the user's own values,
the research photographs, every source with its link (those the user crossed out marked so), every
claim with its evidence status, the open questions, what moves, and the record of how the research
was done (its versions, and every search and page the AI read)."""
from __future__ import annotations

import io
import re
from pathlib import Path
from xml.sax.saxutils import escape

from .jobs import now_iso
from .review import claim_rests_only_on

# A font with more than Latin-1 ("›", "—", "°", curly quotes) where the system has one; reportlab's
# built-in Helvetica can only draw Latin-1, so without one the text is reduced to that.
FONT_CANDIDATES = [
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
    ("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
    ("/Library/Fonts/Arial.ttf", "/Library/Fonts/Arial Bold.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
]
EXCLUDED = "Excluded by you"


def filename(project_name: str, ext: str) -> str:
    stem = re.sub(r"[^A-Za-z0-9]+", "_", project_name).strip("_") or "research"
    return f"{stem}_research.{ext}"


def claim_text(c: dict) -> str:
    """A claim with its value: the research keeps "Length overall" and 362.12 m apart (value, unit)."""
    value = c.get("value")
    if value is None or value == "":
        return str(c.get("claim", ""))
    return f"{c.get('claim', '')}: {value}{' ' + str(c['unit']) if c.get('unit') else ''}"


def _sources_of(c: dict) -> str:
    s = c.get("sources")
    return ", ".join(s) if isinstance(s, list) else str(s or "")


def _fmt(v) -> str:
    if isinstance(v, list):
        return "[" + ", ".join(_fmt(x) for x in v) + "]"
    if isinstance(v, float):
        return f"{v:g}"
    return "—" if v is None else str(v)


def _source_rows(research: dict, excluded: list[str]) -> list[list[str]]:
    return [[s.get("id", ""), f"{s.get('title', '')}\n{s.get('url', '')}",
             f"{EXCLUDED} (was {s.get('status', '')})" if s.get("id") in excluded else s.get("status", ""),
             s.get("what_it_proves", "")] for s in research.get("sources") or []]


def _claim_rows(research: dict, excluded: list[str]) -> list[list[str]]:
    return [[claim_text(c), _sources_of(c), c.get("status", ""),
             f"Not used: its sources were {EXCLUDED.lower()}." if claim_rests_only_on(c, excluded)
             else c.get("consequence_for_model", "")] for c in research.get("claims") or []]


def _plan(project: dict) -> tuple[list[str], list[list[str]]]:
    """The design plan's lines, and the params the user set: param, the plan's value, theirs."""
    plan = project.get("design_plan") or {}
    if not plan:
        return [], []
    lines = [f"Status: {plan.get('status', 'draft')}"] + [str(s) for s in plan.get("summary") or []]
    planned = (plan.get("spec") or {}).get("params") or {}
    mine = plan.get("user_params") or {}
    return lines, [[k, _fmt(planned.get(k)), _fmt(v)] for k, v in mine.items()]


def _record(materials: list[dict], history: list[dict]) -> list[tuple[str, str]]:
    rows = [(f"Research version {h['n']}", f"{h['at']} · {h.get('sources', 0)} sources, {h.get('claims', 0)} claims")
            for h in history]
    rows += [(("Searched: " if m["kind"] == "search" else "Read: ") + m["title"], m["at"]) for m in materials]
    return rows


def _caption(ph: dict) -> str:
    parts = [f"{ph.get('id', '')} {ph.get('title') or ''}".strip(), ph.get("credit") or "", ph.get("date") or "",
             ph.get("source") or ph.get("image_url") or ""]
    return " · ".join(p for p in parts if p)


# ---- PDF -------------------------------------------------------------------------------------

def _pdf_fonts() -> tuple[str, str, bool]:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    for regular, bold in FONT_CANDIDATES:
        if Path(regular).exists() and Path(bold).exists():
            try:
                pdfmetrics.registerFont(TTFont("ReportSans", regular))
                pdfmetrics.registerFont(TTFont("ReportSans-Bold", bold))
                return "ReportSans", "ReportSans-Bold", True
            except Exception:
                continue
    return "Helvetica", "Helvetica-Bold", False


def research_pdf(project: dict, research: dict, materials: list[dict], history: list[dict],
                 photos: list[tuple[dict, bytes]] = ()) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font, bold, unicode_ok = _pdf_fonts()
    excluded = list(project.get("excluded_sources") or [])

    def t(s) -> str:
        s = "" if s is None else str(s)
        if not unicode_ok:
            s = s.encode("latin-1", "replace").decode("latin-1")
        return escape(s).replace("\n", "<br/>")  # a Paragraph ignores plain line breaks

    body = ParagraphStyle("body", fontName=font, fontSize=9.5, leading=13)
    small = ParagraphStyle("small", parent=body, fontSize=8, leading=10.5, textColor=colors.HexColor("#444444"))
    h1 = ParagraphStyle("h1", fontName=bold, fontSize=18, leading=22, spaceAfter=4)
    h2 = ParagraphStyle("h2", fontName=bold, fontSize=12.5, leading=16, spaceBefore=12, spaceAfter=5)
    cell = ParagraphStyle("cell", parent=body, fontSize=8.5, leading=11)
    head = ParagraphStyle("head", parent=cell, fontName=bold)

    def table(header: list[str], rows: list[list[str]], widths: list[float], greyed: set[int] = frozenset()) -> Table:
        data = [[Paragraph(t(h), head) for h in header]] + [[Paragraph(t(c), cell) for c in r] for r in rows]
        tb = Table(data, colWidths=[w * mm for w in widths], repeatRows=1)
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EDF5")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B8C0CC")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ]
        style += [("BACKGROUND", (0, i + 1), (-1, i + 1), colors.HexColor("#F1F1F1")) for i in greyed]
        tb.setStyle(TableStyle(style))
        return tb

    summary = research.get("summary") or {}
    story = [
        Paragraph(t(project["name"]), h1),
        Paragraph(t(research.get("subject")), body),
        Paragraph(t(f"Research report exported {now_iso()} by the 3D model editor. Sources are public; "
                    "evidence statuses are the researcher's judgement."), small),
        Paragraph("Verdict", h2), Paragraph(t(summary.get("verdict")), body),
    ]
    if summary.get("existing_work"):
        story += [Spacer(0, 4), Paragraph(t(f"Existing work: {summary['existing_work']}"), body)]
    lines, mine = _plan(project)
    if lines:
        story += [Paragraph("Design plan", h2)] + [Paragraph(t(line), body) for line in lines]
        if mine:
            story += [Spacer(0, 4), table(["Setting", "The plan's value", "Your value"], mine, [60, 57, 57])]
    if photos:
        story += [Paragraph("Photographs", h2)]
        for ph, raw in photos:
            try:
                w, h = ImageReader(io.BytesIO(raw)).getSize()
            except Exception:
                continue
            width = 120 * mm
            story += [Image(io.BytesIO(raw), width=width, height=width * h / max(w, 1)), Paragraph(t(_caption(ph)), small), Spacer(0, 6)]
    srows = _source_rows(research, excluded)
    story += [Paragraph("Sources", h2), table(["Id", "Source", "Status", "What it proves"], srows, [12, 70, 26, 66],
                                              {i for i, r in enumerate(srows) if r[2].startswith(EXCLUDED)})]
    crows = _claim_rows(research, excluded)
    story += [Paragraph("Claims", h2), table(["Claim", "Sources", "Status", "What it means for the model"], crows,
                                             [56, 20, 26, 72], {i for i, r in enumerate(crows) if r[3].startswith("Not used")})]
    questions = research.get("open_questions") or []
    if questions:
        story += [Paragraph("Open questions", h2), table(
            ["Question", "Meanwhile"],
            [[q.get("question", ""), q.get("meanwhile") or q.get("why_it_matters", "")] for q in questions], [87, 87])]
    moving = research.get("articulation") or []
    if moving:
        story += [Paragraph("What moves", h2), table(
            ["Part", "Motion", "Priority"], [[a.get("part", ""), a.get("motion", ""), a.get("priority", "")] for a in moving],
            [70, 80, 24])]
    record = _record(materials, history)
    if record:
        story += [Paragraph("How this research was done", h2), table(["Step", "When"], [list(r) for r in record], [134, 40])]

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, title=f"{project['name']}: research", leftMargin=16 * mm, rightMargin=16 * mm,
                      topMargin=16 * mm, bottomMargin=16 * mm).build(story)
    return buf.getvalue()


# ---- PowerPoint ------------------------------------------------------------------------------

def research_pptx(project: dict, research: dict, materials: list[dict], history: list[dict],
                  photos: list[tuple[dict, bytes]] = ()) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    excluded = list(project.get("excluded_sources") or [])
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)  # 16:9
    title_layout, blank = deck.slide_layouts[0], deck.slide_layouts[6]

    def slide_with_title(title: str):
        s = deck.slides.add_slide(blank)
        box = s.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.8))
        p = box.text_frame.paragraphs[0]
        p.text, p.font.size, p.font.bold = title, Pt(28), True
        return s

    def bullets(title: str, lines: list[str], size: int = 16):
        s = slide_with_title(title)
        tf = s.shapes.add_textbox(Inches(0.6), Inches(1.3), Inches(12.1), Inches(5.8)).text_frame
        tf.word_wrap = True
        for i, line in enumerate(lines or ["—"]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text, p.font.size, p.space_after = line, Pt(size), Pt(6)

    def table_slides(title: str, header: list[str], rows: list[list[str]], widths: list[float], per_slide: int = 7):
        for start in range(0, max(len(rows), 1), per_slide):
            chunk = rows[start:start + per_slide]
            s = slide_with_title(title + (f" ({start // per_slide + 1})" if len(rows) > per_slide else ""))
            shape = s.shapes.add_table(len(chunk) + 1, len(header), Inches(0.6), Inches(1.3), Inches(12.1), Inches(0.4))
            tb = shape.table
            for j, w in enumerate(widths):
                tb.columns[j].width = Inches(w)
            for j, h in enumerate(header):
                tb.cell(0, j).text = h
            for i, row in enumerate(chunk, start=1):
                for j, value in enumerate(row):
                    tb.cell(i, j).text = str(value)
            for r in range(len(chunk) + 1):
                for j in range(len(header)):
                    for p in tb.cell(r, j).text_frame.paragraphs:
                        p.font.size = Pt(12)

    cover = deck.slides.add_slide(title_layout)
    cover.shapes.title.text = project["name"]
    cover.placeholders[1].text = f"{research.get('subject', '')}\nResearch report · {now_iso()[:10]}"

    summary = research.get("summary") or {}
    bullets("Verdict", [summary.get("verdict") or "—"] + ([f"Existing work: {summary['existing_work']}"]
                                                         if summary.get("existing_work") else []), 20)
    lines, mine = _plan(project)
    if lines:
        bullets("Design plan", lines, 15)
        if mine:
            table_slides("Design plan: your values", ["Setting", "The plan's value", "Your value"], mine, [4.0, 4.0, 4.1], per_slide=9)
    for i in range(0, len(photos), 2):
        s = slide_with_title("Photographs" + (f" ({i // 2 + 1})" if len(photos) > 2 else ""))
        for j, (ph, raw) in enumerate(photos[i:i + 2]):
            left = Inches(0.6 + j * 6.2)
            try:
                s.shapes.add_picture(io.BytesIO(raw), left, Inches(1.3), width=Inches(5.9))
            except Exception:
                continue
            cap = s.shapes.add_textbox(left, Inches(6.5), Inches(5.9), Inches(0.8)).text_frame
            cap.word_wrap = True
            cap.paragraphs[0].text, cap.paragraphs[0].font.size = _caption(ph), Pt(11)
    table_slides("Sources", ["Id", "Source", "Status", "What it proves"], _source_rows(research, excluded), [0.8, 5.0, 1.8, 4.5])
    table_slides("Claims", ["Claim", "Sources", "Status", "What it means for the model"], _claim_rows(research, excluded),
                 [3.8, 1.3, 1.8, 5.2])
    bullets("Open questions", [f"{q.get('question', '')}: {q.get('meanwhile') or q.get('why_it_matters') or ''}".rstrip(": ")
                               for q in research.get("open_questions") or []])
    if research.get("articulation"):
        table_slides("What moves", ["Part", "Motion", "Priority"],
                     [[a.get("part", ""), a.get("motion", ""), a.get("priority", "")] for a in research["articulation"]],
                     [5.0, 5.6, 1.5])
    record = _record(materials, history)
    if record:
        table_slides("How this research was done", ["Step", "When"], [list(r) for r in record], [9.4, 2.7], per_slide=10)

    buf = io.BytesIO()
    deck.save(buf)
    return buf.getvalue()
